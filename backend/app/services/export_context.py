"""Colunas de contexto comuns às três planilhas (card #1123, decisão 7).

Um único construtor para a planilha de Favoritos, a de trades do resultado e a do ranking da
Descoberta: os 22 campos na ordem do spec ``result-context-export``, cada valor com o estado
``registrado | deduzido | não registrado``. "deduzido" vira sufixo "(deduzido)"; "não registrado" é
o próprio texto. Números saem como número com formato (% ou p.p.); célula sem comparação fica vazia,
nunca zero. Nada aqui coleta dado ou recalcula comparação: tudo vem do que foi gravado.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Iterable, Optional

from app.services.benchmark_comparison import (
    CLASS_LABEL,
    CURRENCY_NOTE_TEXT,
    MIN_COMMON_DAYS,
    DAYS_PER_YEAR,
    br_date,
    return_kind_of,
    return_kind_text,
    to_day,
)

REGISTRADO = "registrado"
DEDUZIDO = "deduzido"
NAO_REGISTRADO = "não registrado"

CONTEXT_FIELDS = (
    "Classe",
    "Moeda",
    "Fonte",
    "Sessão",
    "Ajuste",
    "Período pedido",
    "Período efetivo",
    "Modo de execução",
    "Versão da amostra",
    "Referência",
    "Tipo de retorno da referência",
    "Período comum",
    "Fora do período comum",
    "Retorno da estratégia no período comum",
    "CAGR da estratégia no período comum",
    "Retorno da referência",
    "CAGR da referência",
    "Diferença de CAGR (p.p.)",
    "Situação da comparação",
    "Aviso de moeda",
    "Versão da amostra da referência",
    "Comparação gravada em",
)
COMPARISON_FIELDS = CONTEXT_FIELDS[9:]
EXTRA_FIELDS = tuple(f"{name} (extra)" for name in COMPARISON_FIELDS)

PCT_FORMAT = "+0.00%;-0.00%;0.00%"
PP_FORMAT = "+0.0;-0.0;0.0"
DISCOVERY_REFERENCE = "Próprio ativo (buy-and-hold, long-only)"

LEGEND_BASE = (
    ("deduzido", "Não foi gravado no favorito; saiu do símbolo (classe e moeda)."),
    (
        "não registrado",
        "Não foi gravado e não dá para deduzir (fonte, sessão, ajuste, versão da amostra, comparação).",
    ),
    (
        "USDT tratado como USD",
        "Estratégia e referência em moedas diferentes: comparação só em percentual, sem conversão.",
    ),
    (
        "Retorno de preço, sem dividendos",
        "Referência de ação ou ETF medida só pelo preço ajustado por split.",
    ),
    (
        "Período comum",
        f"Só o período que estratégia e referência cobrem, em fechamentos diários; CAGR com dias "
        f"corridos ÷ 365,25. Abaixo de {MIN_COMMON_DAYS} dias não há comparação.",
    ),
)
LEGEND_REDONE = (
    "Comparação refeita",
    "Só quando o backtest do favorito é atualizado, com a mesma referência. Abrir ou exportar não recalcula.",
)
LEGEND_SCREEN = (
    "Tela × planilha",
    "A tela de Favoritos mostra a convenção da classe (fonte e sessão). A planilha só traz o que foi "
    "gravado e marca o resto como “não registrado”.",
)
LEGEND_RANKING = (
    "Retornos do ranking",
    "Derivados do CAGR e da janela gravados na varredura, sem dado novo. Nos 22 campos de contexto, "
    "a referência é sempre o próprio ativo.",
)
LEGEND_EXTRA = (
    "(extra)",
    "Colunas da referência extra escolhida no Decidir, com os mesmos nomes dos campos de contexto. "
    "Não mudam a ordem, o rank global nem a elegibilidade; o resto da linha segue o próprio ativo, "
    "por isso o retorno contra o próprio ativo e o (extra) diferem na mesma linha.",
)


@dataclass
class Cell:
    value: Any
    state: str = REGISTRADO
    fmt: Optional[str] = None

    def rendered(self) -> Any:
        if self.state == NAO_REGISTRADO:
            return NAO_REGISTRADO
        if self.state == DEDUZIDO and isinstance(self.value, str) and self.value:
            return f"{self.value} (deduzido)"
        return self.value


def nr() -> Cell:
    return Cell(NAO_REGISTRADO, NAO_REGISTRADO)


def empty() -> Cell:
    return Cell("", REGISTRADO)


def _num(value: Any) -> Optional[float]:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def pct(value: Any) -> Cell:
    number = _num(value)
    return Cell(number, REGISTRADO, PCT_FORMAT) if number is not None else empty()


def pp(value: Any) -> Cell:
    number = _num(value)
    return Cell(round(number, 4), REGISTRADO, PP_FORMAT) if number is not None else empty()


# --------------------------------------------------------------------------- rótulos


def class_label(asset_class: Optional[str]) -> Optional[str]:
    return CLASS_LABEL.get(str(asset_class or "").lower())


def deduce_class(symbol: str) -> tuple[str, str]:
    """Regra da tela: par com ``/`` ⇒ Cripto e a cotação do par; ticker ⇒ catálogo (Ação ou ETF) e USD."""

    sym = str(symbol or "").upper()
    if "/" in sym:
        return "crypto", sym.split("/", 1)[1] or "USDT"
    try:
        from app.services.equity_backtest import asset_identity

        ident = asset_identity(sym)
        return ident.get("asset_class") or "stock", ident.get("currency") or "USD"
    except Exception:
        return "stock", "USD"


_SOURCE = {
    "ccxt": "Binance Spot",
    "binance_spot": "Binance Spot",
    "binance spot": "Binance Spot",
    "binance_equity": "Binance Ações",
    "binance ações": "Binance Ações",
    "stooq": "Stooq · fonte antiga",
}
_SESSION = {
    "24x7": "24/7",
    "24/7": "24/7",
    "rth": "Pregão regular de Nova York",
    "regular_ny": "Pregão regular de Nova York",
    "eod": "Diário · fonte antiga",
}
_ADJUSTMENT = {
    "na": "Não se aplica",
    "none": "Não se aplica",
    "adjusted": "Ajustado por split",
    "split": "Ajustado por split",
    "raw": "Sem ajuste",
    "legacy": "Fonte antiga",
}


def source_label(value: Any) -> Optional[str]:
    text = str(value or "").strip()
    return _SOURCE.get(text.lower(), text or None)


def session_label(value: Any) -> Optional[str]:
    text = str(value or "").strip()
    return _SESSION.get(text.lower(), text or None)


def adjustment_label(value: Any) -> Optional[str]:
    text = str(value or "").strip()
    return _ADJUSTMENT.get(text.lower(), text or None)


def period_text(start: Any, end: Any, *, all_label: str = "Todo o histórico") -> Optional[str]:
    a, b = to_day(start), to_day(end)
    if a and b:
        return f"{br_date(a)} – {br_date(b)}"
    if a:
        return f"desde {br_date(a)}"
    if b:
        return f"até {br_date(b)}"
    return all_label


def window_text(
    start: Any, end: Any, asset_class: Optional[str] = None, candles: Any = None
) -> Optional[str]:
    a, b = to_day(start, asset_class), to_day(end, asset_class)
    if not (a and b):
        return None
    text = f"{br_date(a)} – {br_date(b)}"
    if candles:
        text += f" · {int(candles):,} candles".replace(",", ".")
    return text


def precision_text(
    execution: Optional[dict[str, Any]], mode: Optional[str] = None
) -> Optional[str]:
    if isinstance(execution, dict) and execution:
        if execution.get("precision_effective") == "intrabar_15m":
            return "intrabar 15m"
        if execution.get("precision_requested") == "intrabar":
            return "precisão diária (intrabar sem cobertura)"
        return "precisão do candle"
    if mode == "deep_15m":
        return "intrabar 15m"
    if mode == "fast_1d":
        return "precisão do candle"
    return None


def sample_text(
    provenance: Optional[dict[str, Any]],
    symbol: str,
    timeframe: Optional[str],
    snapshot_id: Optional[str],
) -> Optional[str]:
    prov = provenance if isinstance(provenance, dict) else {}
    version = prov.get("version_label") or prov.get("data_version")
    sid = snapshot_id or prov.get("snapshot_id")
    parts = []
    if version not in (None, ""):
        parts.append(f"{symbol} {timeframe or ''} versão {version}".replace("  ", " ").strip())
    if sid:
        parts.append(str(sid))
    return " · ".join(parts) if parts else None


def when_text(value: Any) -> Optional[str]:
    if not value:
        return None
    try:
        import pandas as pd

        ts = pd.Timestamp(value)
        if ts.tzinfo is None:
            ts = ts.tz_localize("UTC")
        return ts.tz_convert("America/Sao_Paulo").strftime("%d/%m/%Y %H:%M")
    except Exception:
        return str(value)


# --------------------------------------------------------------------------- comparação (13 campos)


def reference_text(payload: dict[str, Any]) -> str:
    ref = payload.get("reference") or {}
    symbol = ref.get("symbol") or "—"
    if ref.get("kind") == "own":
        return f"{symbol} (próprio ativo)"
    parts = [symbol, class_label(ref.get("asset_class")) or "", ref.get("currency") or ""]
    return " · ".join(p for p in parts if p)


def common_text(common: Optional[dict[str, Any]]) -> str:
    if not common:
        return ""
    a, b = to_day(common.get("start")), to_day(common.get("end"))
    days = int(common.get("days") or 0)
    return f"{br_date(a)} – {br_date(b)} · {days:,} dias".replace(",", ".")


def left_out_text(payload: dict[str, Any]) -> str:
    parts = payload.get("left_out") or []
    if not parts:
        return "nada"
    out = []
    for part in parts:
        a, b = to_day(part.get("start")), to_day(part.get("end"))
        span = br_date(a) if a == b else f"{br_date(a)} – {br_date(b)}"
        if part.get("note") == "no_session":
            span += " (sem pregão em Nova York)"
        trades = int(part.get("trades") or 0)
        out.append(f"{span} · {trades} {'trade' if trades == 1 else 'trades'}")
    return "; ".join(out)


def reference_sample_text(payload: dict[str, Any]) -> Optional[str]:
    sample = ((payload.get("samples") or {}).get("reference")) or {}
    symbol = (payload.get("reference") or {}).get("symbol") or ""
    version = sample.get("version_label")
    if version in (None, "") and not sample.get("snapshot_id"):
        return None
    parts = [f"{symbol} versão {version}" if version not in (None, "") else symbol]
    if sample.get("snapshot_id"):
        parts.append(str(sample["snapshot_id"]))
    return " · ".join(parts)


def situation_text(payload: dict[str, Any], *, short: bool = False) -> str:
    if payload.get("status") == "valid":
        return "Válida"
    reason = payload.get("reason") or {}
    text = (
        (reason.get("short") if short else reason.get("text")) or reason.get("text") or "sem motivo"
    )
    long_text = str(reason.get("text") or "")
    if reason.get("code") == "sample_unavailable" or (
        # Cripto sem amostra cujo refeito falhou: o texto longo já diz "Sem comparação:".
        reason.get("code") == "regeneration_failed"
        and long_text.startswith("Sem comparação")
    ):
        # Card #1123 (retrabalho): o histórico não é insuficiente; os trades da varredura é que não
        # puderam ser refeitos iguais (conferidos pelas métricas). O texto longo já começa com
        # "Sem comparação:".
        return text if text.startswith("Sem comparação") else f"Sem comparação: {text}"
    return f"Histórico insuficiente: {text}"


def comparison_cells(
    payload: Optional[dict[str, Any]], *, short_reason: bool = False
) -> list[Cell]:
    """Os 13 campos da referência (de Referência a Comparação gravada em)."""

    if not payload:
        return [Cell("Próprio ativo")] + [empty()] * 8 + [nr(), empty(), nr(), nr()]
    valid = payload.get("status") == "valid"
    ref = payload.get("reference") or {}
    strategy = payload.get("strategy") or {}
    ref_result = payload.get("reference_result") or {}
    delta = payload.get("delta") or {}
    sample = reference_sample_text(payload)
    return [
        Cell(reference_text(payload)),
        # O tipo de retorno vale para a referência escolhida, mesmo sem comparação (P3 aceito).
        Cell(return_kind_text(ref.get("return_kind") or return_kind_of(ref.get("asset_class")))),
        Cell(common_text(payload.get("common_window")) if valid else ""),
        Cell(left_out_text(payload) if valid else ""),
        pct(strategy.get("return")) if valid else empty(),
        pct(strategy.get("cagr")) if valid else empty(),
        pct(ref_result.get("return")) if valid else empty(),
        pct(ref_result.get("cagr")) if valid else empty(),
        pp(delta.get("cagr_pp")) if valid else empty(),
        Cell(situation_text(payload, short=short_reason)),
        Cell(CURRENCY_NOTE_TEXT if valid and payload.get("currency_note") else ""),
        Cell(sample) if sample else nr(),
        Cell(when_text(payload.get("created_at")) or "") if payload.get("created_at") else nr(),
    ]


# --------------------------------------------------------------------------- 9 campos do resultado


def identity_cells(
    *,
    symbol: str,
    asset: Optional[dict[str, Any]],
    asset_recorded: bool,
    source: Any,
    session: Any,
    adjustment: Any,
    provenance_recorded: bool,
    requested: Optional[str],
    effective: Optional[str],
    mode: Optional[str],
    sample: Optional[str],
) -> list[Cell]:
    asset = asset or {}
    if asset_recorded and asset.get("asset_class"):
        cls = Cell(class_label(asset.get("asset_class")) or str(asset.get("asset_class")))
        cur = Cell(asset.get("currency") or ("USDT" if "/" in symbol else "USD"))
    else:
        deduced_class, deduced_currency = deduce_class(symbol)
        cls = Cell(class_label(asset.get("asset_class") or deduced_class), DEDUZIDO)
        cur = Cell(asset.get("currency") or deduced_currency, DEDUZIDO)

    def recorded(value: Optional[str]) -> Cell:
        return Cell(value) if provenance_recorded and value else nr()

    return [
        cls,
        cur,
        recorded(source_label(source)),
        recorded(session_label(session)),
        recorded(adjustment_label(adjustment)),
        Cell(requested) if requested else nr(),
        Cell(effective) if effective else nr(),
        Cell(mode) if mode else nr(),
        Cell(sample) if sample else nr(),
    ]


def result_context(context: dict[str, Any], comparison: Optional[dict[str, Any]]) -> list[Cell]:
    """Contexto enviado pela tela do resultado (rodada nova): asset, data_provenance, execution…"""

    symbol = str(context.get("symbol") or "")
    asset = context.get("asset") if isinstance(context.get("asset"), dict) else None
    prov = (
        context.get("data_provenance") if isinstance(context.get("data_provenance"), dict) else None
    )
    execution = context.get("execution") if isinstance(context.get("execution"), dict) else None
    window = ((execution or {}).get("window") or {}).get("covered") or {}
    is_crypto = "/" in symbol
    source = (prov or {}).get("data_source") or (
        "Binance Spot" if is_crypto and execution else None
    )
    session = (prov or {}).get("session") or ("24x7" if is_crypto and execution else None)
    adjustment = (prov or {}).get("adjustment") or ("na" if is_crypto and execution else None)
    kind = "Otimização" if context.get("is_optimization") else "Backtest"
    precision = precision_text(execution, context.get("execution_mode"))
    mode = f"{kind} · {precision}" if precision else None
    return identity_cells(
        symbol=symbol,
        asset=asset,
        asset_recorded=bool(asset and asset.get("asset_class")),
        source=source,
        session=session,
        adjustment=adjustment,
        provenance_recorded=bool(prov or (is_crypto and execution)),
        requested=(
            period_text(context.get("start_date"), context.get("end_date"))
            if (
                context.get("start_date")
                or context.get("end_date")
                or context.get("period_type") == "all"
            )
            else None
        ),
        effective=window_text(
            window.get("start"),
            window.get("end"),
            (asset or {}).get("asset_class"),
            window.get("candles"),
        ),
        mode=mode,
        sample=sample_text(prov, symbol, context.get("timeframe"), context.get("snapshot_id")),
    ) + comparison_cells(comparison)


def favorite_context(
    fav: Any, asset: Optional[dict[str, Any]], comparison: Optional[dict[str, Any]]
) -> list[Cell]:
    metrics = fav.metrics if isinstance(fav.metrics, dict) else {}
    symbol = str(fav.symbol or "")
    stored_asset = metrics.get("analysis_asset") or metrics.get("asset")
    asset_recorded = isinstance(stored_asset, dict) and bool(stored_asset.get("asset_class"))
    prov = (
        metrics.get("data_provenance") if isinstance(metrics.get("data_provenance"), dict) else None
    )
    execution = (
        metrics.get("analysis_execution")
        if isinstance(metrics.get("analysis_execution"), dict)
        else None
    )
    provenance_recorded = bool(prov or execution or fav.snapshot_id)
    window = ((execution or {}).get("window") or {}).get("covered") or {}
    precision = precision_text(execution, metrics.get("analysis_execution_mode"))
    requested = (
        period_text(fav.start_date, fav.end_date)
        if (fav.start_date or fav.end_date or fav.period_type == "all")
        else None
    )
    return identity_cells(
        symbol=symbol,
        asset=stored_asset if asset_recorded else (asset or {}),
        asset_recorded=asset_recorded,
        source=(prov or {}).get("data_source") or fav.data_source,
        session=(prov or {}).get("session") or fav.session,
        adjustment=(prov or {}).get("adjustment") or fav.adjustment,
        provenance_recorded=provenance_recorded,
        requested=requested,
        effective=window_text(
            window.get("start"),
            window.get("end"),
            (stored_asset or asset or {}).get("asset_class"),
            window.get("candles"),
        ),
        mode=f"Favorito · {precision}" if precision else None,
        sample=(
            sample_text(prov, symbol, fav.timeframe, fav.snapshot_id)
            if (prov or fav.snapshot_id)
            else None
        ),
    ) + comparison_cells(comparison)


def derived_return(cagr: Any, days: int) -> Optional[float]:
    value = _num(cagr)
    if value is None or days <= 0:
        return None
    base = 1.0 + value
    if base <= 0:
        return -1.0
    return base ** (days / DAYS_PER_YEAR) - 1.0


def discovery_context(row: dict[str, Any], sweep: dict[str, Any]) -> list[Cell]:
    """22 campos de um candidato do ranking: referência = próprio ativo, sem dado novo."""

    asset = row.get("asset") or {}
    prov = row.get("provenance") or {}
    asset_class = asset.get("asset_class")
    start, end = row.get("start_at"), row.get("end_at")
    a, b = to_day(start, asset_class), to_day(end, asset_class)
    days = (b - a).days if a and b else 0
    precision = precision_text(row.get("execution"))
    mode = f"Varredura · {precision}" if precision else "Varredura"
    sample = (
        " · ".join(str(p) for p in (row.get("candle_version"), row.get("snapshot_id")) if p) or None
    )
    completed = when_text(sweep.get("completed_at"))
    recorded_at = (
        f"{completed} · varredura #{sweep.get('sweep_id') or row.get('sweep_id')}"
        if completed
        else nr()
    )
    cells = identity_cells(
        symbol=str(row.get("symbol") or ""),
        asset=asset,
        asset_recorded=bool(asset_class),
        source=prov.get("data_source"),
        session=prov.get("session"),
        adjustment=prov.get("adjustment"),
        provenance_recorded=bool(prov.get("data_source")),
        requested=period_text(sweep.get("start_date"), sweep.get("end_date")),
        effective=window_text(start, end, asset_class),
        mode=mode,
        sample=sample,
    )
    cells.append(Cell(DISCOVERY_REFERENCE))
    cagr, bh = _num(row.get("cagr")), _num(row.get("benchmark_cagr"))
    if row.get("eligibility") == "insufficient_sample" or cagr is None:
        situation = "Sem amostra: o candidato não tem métrica gravada."
        cells += [empty()] * 8 + [
            Cell(situation),
            empty(),
            Cell("a mesma do candidato"),
            Cell(recorded_at) if isinstance(recorded_at, str) else recorded_at,
        ]
        return cells
    if days < MIN_COMMON_DAYS:
        cells += (
            [Cell(return_kind_text(return_kind_of(asset_class)))]
            + [empty()] * 7
            + [
                Cell(
                    f"Histórico insuficiente: Período comum de {days} dias: o mínimo é {MIN_COMMON_DAYS}."
                ),
                empty(),
                Cell("a mesma do candidato"),
                Cell(recorded_at) if isinstance(recorded_at, str) else recorded_at,
            ]
        )
        return cells
    cells += [
        Cell(return_kind_text(return_kind_of(asset_class))),
        Cell(f"{br_date(a)} – {br_date(b)} · {days:,} dias".replace(",", ".")),
        Cell("nada"),
        pct(derived_return(cagr, days)),
        pct(cagr),
        pct(derived_return(bh, days)) if bh is not None else empty(),
        pct(bh) if bh is not None else empty(),
        pp((cagr - bh) * 100.0) if bh is not None else empty(),
        Cell("Válida"),
        empty(),
        Cell("a mesma do candidato"),
        Cell(recorded_at) if isinstance(recorded_at, str) else recorded_at,
    ]
    return cells


# --------------------------------------------------------------------------- legenda e escrita


def legend_rows(
    *,
    favorite: bool = False,
    favorites_screen: bool = False,
    ranking: bool = False,
    extra: bool = False,
) -> list[tuple[str, str]]:
    rows = list(LEGEND_BASE)
    if favorites_screen:
        rows.append(LEGEND_SCREEN)
    if favorite or favorites_screen:
        rows.append(LEGEND_REDONE)
    if ranking:
        rows.append(LEGEND_RANKING)
    if extra:
        rows.append(LEGEND_EXTRA)
    return rows


def write_workbook(sheets: Iterable[tuple[str, list[str], list[list[Any]]]]) -> bytes:
    """Abas na ordem dada; ``Cell`` vira valor com formato, o resto entra como está."""

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        ws.append(list(header))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        widths = [len(str(h)) for h in header]
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, item in enumerate(row, start=1):
                value = item.rendered() if isinstance(item, Cell) else item
                target = ws.cell(row=r_idx, column=c_idx, value=value)
                if isinstance(item, Cell) and item.fmt and isinstance(item.value, (int, float)):
                    target.number_format = item.fmt
                length = len(str(value)) if value is not None else 0
                if c_idx - 1 < len(widths):
                    widths[c_idx - 1] = max(widths[c_idx - 1], min(length, 60))
        for idx, width in enumerate(widths, start=1):
            from openpyxl.utils import get_column_letter

            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 62)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def legend_sheet(**kwargs: bool) -> tuple[str, list[str], list[list[Any]]]:
    return ("Legenda", ["Termo", "Significado"], [list(r) for r in legend_rows(**kwargs)])


def context_sheet(cells: list[Cell]) -> tuple[str, list[str], list[list[Any]]]:
    rows = []
    for name, cell in zip(CONTEXT_FIELDS, cells):
        rows.append([name, cell, cell.state])
    return ("Contexto", ["Campo", "Valor", "Estado"], rows)


def today_iso() -> str:
    return datetime.now(timezone.utc).date().isoformat()


__all__ = [
    "CONTEXT_FIELDS",
    "COMPARISON_FIELDS",
    "EXTRA_FIELDS",
    "Cell",
    "comparison_cells",
    "context_sheet",
    "discovery_context",
    "favorite_context",
    "legend_sheet",
    "result_context",
    "write_workbook",
]
