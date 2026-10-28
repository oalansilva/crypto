"""Card #1123: coluna extra do Decidir (escolha gravada por varredura e usuário, ranking intocado)."""

from __future__ import annotations

import io
from datetime import datetime, timezone

import openpyxl
import pandas as pd
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models_benchmark import BenchmarkComparison, DiscoveryReferenceColumn
from app.routes import discovery_routes
from app.services import benchmark_comparison as bench
from app.services import discovery_reference_column as refcol
from app.services import export_context as xc
from support_card_1123 import (
    sync_e2e_fixture,
    BTC_TRADES,
    FakeSource,
    client_for,
    ensure_snapshot_table,
    insert_snapshot,
    make_user,
    pinned,
    trade,
)

SWEEP = "c1123mix"

CANDIDATES = [
    # id, símbolo, classe, tf, início, fim, cagr, b&h, calmar, elegibilidade
    (
        "RS-BTC1123A1",
        "BTC/USDT",
        "crypto",
        "1d",
        "2019-01-01",
        "2024-12-31",
        0.82,
        0.7057,
        4.1,
        "eligible",
    ),
    (
        "RS-NVDA1123B2",
        "NVDA",
        "stock",
        "1d",
        "2019-01-02",
        "2024-12-31",
        0.65,
        0.848,
        3.0,
        "eligible",
    ),
    (
        "RS-SOL1123D4",
        "SOL/USDT",
        "crypto",
        "1d",
        "2020-08-11",
        "2024-12-31",
        0.95,
        1.5872,
        2.0,
        "eligible",
    ),
    ("RS-SPY1123C3", "SPY", "etf", "1d", "2019-01-02", "2024-12-31", 0.12, 0.1525, 1.0, "eligible"),
    (
        "RS-PENGU1123E5",
        "PENGU/USDT",
        "crypto",
        "1h",
        "2024-12-17",
        "2024-12-31",
        0.40,
        0.30,
        0.5,
        "low_sample",
    ),
    (
        "RS-AAPL1123F6",
        "AAPL",
        "stock",
        "1d",
        "2019-01-02",
        "2024-12-31",
        None,
        None,
        None,
        "insufficient_sample",
    ),
]

TRADES = {
    "RS-BTC1123A1": BTC_TRADES,
    "RS-NVDA1123B2": [
        trade("2019-03-01", "2019-09-03", 0.4),
        trade("2021-01-04", "2021-06-01", 0.3),
    ],
    "RS-SOL1123D4": [trade("2020-09-01", "2021-04-01", 1.2)],
    "RS-SPY1123C3": [trade("2019-03-01", "2019-09-03", 0.05)],
    "RS-PENGU1123E5": [trade("2024-12-18", "2024-12-20", 0.1)],
}


def _seed(db, sweep_id=SWEEP, actor="admin-1", expired: set[str] | None = None):
    from app.models_discovery import DiscoveryResult, DiscoverySweep

    db.add(
        DiscoverySweep(
            id=sweep_id,
            actor=actor,
            state="completed",
            idempotency_key=f"key-{sweep_id}",
            payload_hash="a" * 64,
            snapshot_token=f"tok-{sweep_id}",
            snapshot_hash="b" * 64,
            snapshot={"start_date": "2019-01-01", "end_date": "2024-12-31"},
            total=len(CANDIDATES),
            completed_at=datetime(2026, 10, 27, 12, 0),
        )
    )
    for n, (rid, sym, cls, tf, start, end, cagr, bh, calmar, elig) in enumerate(CANDIDATES):
        rid = rid if sweep_id == SWEEP else f"{rid}-{sweep_id}"
        sid = f"snap{n}{sweep_id}"[:40]
        past = expired and rid in expired
        insert_snapshot(
            db,
            sid,
            retain_until=pd.Timestamp("2026-09-30", tz="UTC") if past else None,
        )
        db.add(
            DiscoveryResult(
                id=rid,
                sweep_id=sweep_id,
                combination_id=n + (0 if sweep_id == SWEEP else 100),
                template_id="multi_ma_crossover",
                symbol=sym,
                timeframe=tf,
                direction="long",
                parameters={"direction": "long", "fast": 10},
                start_at=datetime.fromisoformat(start),
                end_at=datetime.fromisoformat(end),
                metrics={"oos_verdict": {"status": "GO"}} if cagr is not None else {},
                trades_count=40 if cagr is not None else 0,
                calmar_ratio=calmar,
                cagr=cagr,
                benchmark_cagr=bh,
                # Mesma unidade do job (``discovery_tasks._ranking_columns``): p.p.
                delta_cagr_vs_bh=(cagr - bh) * 100 if cagr is not None else None,
                coverage=1.0,
                strategy_identity_key=f"id-{rid}",
                evidence_fingerprint=f"fp-{rid}",
                eligibility=elig,
                dedup_state="unique",
                asset_class=cls,
                data_source="Binance Spot" if cls == "crypto" else "Binance Ações",
                session="24x7" if cls == "crypto" else "rth",
                adjustment="na" if cls == "crypto" else "adjusted",
                snapshot_id=sid,
                candle_version="v1",
            )
        )
    db.commit()


@pytest.fixture
def db(postgres_isolation, unit_database_url):
    engine = create_engine(unit_database_url)
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = Session()
    ensure_snapshot_table(session)
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture
def source():
    fake = FakeSource()
    bench.set_reference_source(fake)
    yield fake
    bench.set_reference_source(None)


@pytest.fixture
def ctx(db, source, monkeypatch):
    user = make_user(db, "alan", admin=True)
    actor = str(user.id)
    _seed(db, actor=actor)
    state = {"user": actor, "admin": True}
    queued: list[str] = []
    monkeypatch.setattr(refcol, "enqueue", lambda cid: queued.append(cid))
    regenerated: list[str] = []

    def regenerate(_db, result):
        regenerated.append(result.id)
        return TRADES.get(result.id, [])

    Session = sessionmaker(bind=db.get_bind())

    def run(column_id=None):
        cid = column_id or queued.pop(0)
        refcol.run_column(cid, db_factory=Session, source=source, regenerate=regenerate)
        return cid

    client = client_for(db, [discovery_routes.router], state)
    return {
        "client": client,
        "actor": actor,
        "run": run,
        "queued": queued,
        "regenerated": regenerated,
        "state": state,
    }


def _board(client, metric="delta_cagr_vs_bh", **params):
    res = client.get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard", params={"metric": metric, **params}
    )
    assert res.status_code == 200, res.text
    return res.json()


def _order(board):
    return [(r["result_id"], r["rank"]) for r in board["results"]]


def _choose(client, symbol):
    return client.put(
        f"/api/combos/discovery/sweeps/{SWEEP}/reference-column", json={"reference_symbol": symbol}
    )


def _cells(board):
    return {r["result_id"]: r.get("reference_comparison") for r in board["results"]}


def test_spy_column_matches_the_result_and_keeps_order_and_rank(ctx, source):
    client = ctx["client"]
    before = {m: _order(_board(client, m)) for m in ("delta_cagr_vs_bh", "calmar_ratio")}
    res = _choose(client, "SPY")
    assert res.status_code == 202
    assert res.json()["pending"]["status"] == "computing"
    pending = _board(client)
    assert pending["reference_column"]["column"] is None  # ainda não virou a coluna
    ctx["run"]()
    board = _board(client)
    col = board["reference_column"]["column"]
    assert col["status"] == "done" and col["reference"]["symbol"] == "SPY"
    assert col["counts"]["valid"] == 4 and col["counts"]["short_overlap"] == 1
    assert col["counts"]["currency_note"] is True
    for metric in before:
        assert _order(_board(client, metric)) == before[metric]
    cells = _cells(board)
    btc = cells["RS-BTC1123A1"]
    assert btc["state"] == "valid"
    assert btc["common_window"] == {"start": "2019-01-02", "end": "2024-12-31", "days": 2190}
    assert btc["currency_note"] == "usdt_as_usd" and btc["trades_out"] == 0
    # Mesmo número da seção do resultado (mesma função e mesma série).
    same = bench.compute(
        bench.strategy_side(
            symbol="BTC/USDT",
            asset=None,
            window={"start": "2019-01-01", "end": "2024-12-31"},
            trades=BTC_TRADES,
        ),
        {**source.identity("SPY"), "kind": "catalog"},
        bench.daily_closes(source.frames["SPY"], "etf"),
    )
    assert btc["delta_cagr_pp"] == pytest.approx(same["delta"]["cagr_pp"])
    assert cells["RS-NVDA1123B2"]["currency_note"] is None
    assert cells["RS-PENGU1123E5"] == {
        "state": "insufficient",
        "id": cells["RS-PENGU1123E5"]["id"],
        "reason_code": "short_overlap",
        "reason": "Período comum de 14 dias: o mínimo é 30.",
        "expired_at": None,
    }
    assert cells["RS-AAPL1123F6"] == {"state": "no_sample"}
    assert "RS-AAPL1123F6" not in ctx["regenerated"]


def test_reading_and_exporting_never_compute_or_collect(ctx, source, monkeypatch):
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    loads, regenerated = list(source.loads), list(ctx["regenerated"])

    def boom(*_a, **_k):
        raise AssertionError("leitura não calcula")

    monkeypatch.setattr(bench, "compute", boom)
    monkeypatch.setattr(bench, "build_payload", boom)
    _board(client)
    _board(client)  # recarregar
    res = client.get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard/export",
        params={"metric": "delta_cagr_vs_bh"},
    )
    assert res.status_code == 200
    assert source.loads == loads and ctx["regenerated"] == regenerated


def test_choosing_again_a_recorded_reference_reactivates_without_computing(ctx, source, db):
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    first = _board(client)["reference_column"]["column"]["id"]
    assert _choose(client, None).status_code == 200
    assert _board(client)["reference_column"] is None
    loads = list(source.loads)
    res = _choose(client, "SPY")
    assert res.status_code == 200 and res.json()["reactivated"] is True
    assert ctx["queued"] == []
    assert source.loads == loads
    assert _board(client)["reference_column"]["column"]["id"] == first
    assert db.query(DiscoveryReferenceColumn).count() == 1


def test_crcl_has_no_common_period_with_any_candidate(ctx):
    client = ctx["client"]
    _choose(client, "CRCL")
    ctx["run"]()
    board = _board(client)
    cells = _cells(board)
    for rid, cell in cells.items():
        if rid == "RS-AAPL1123F6":
            continue
        assert cell["state"] == "insufficient" and cell["reason_code"] == "no_overlap"
        assert cell["reason"] == "Sem período comum com CRCL."
        assert "delta_cagr_pp" not in cell
    assert board["reference_column"]["column"]["counts"]["no_overlap"] == 5
    assert board["reference_column"]["column"]["counts"]["valid"] == 0


def test_collection_failure_records_nothing_keeps_the_previous_column_and_retries(ctx, source, db):
    client = ctx["client"]
    _choose(client, "CRCL")
    ctx["run"]()
    cells_before = db.query(BenchmarkComparison).count()
    source.fail.add("QQQ")
    source.missing.add("QQQ")
    res = _choose(client, "QQQ")
    assert res.status_code == 202 and res.json()["pending"]["status"] == "collecting"
    ctx["run"]()
    board = _board(client)
    assert board["reference_column"]["column"]["reference"]["symbol"] == "CRCL"
    assert board["reference_column"]["pending"]["status"] == "failed"
    assert "tempo esgotado" in board["reference_column"]["pending"]["error"]
    assert db.query(BenchmarkComparison).count() == cells_before
    source.fail.discard("QQQ")
    retry = client.post(f"/api/combos/discovery/sweeps/{SWEEP}/reference-column/retry")
    assert retry.status_code == 202
    ctx["run"]()
    after = _board(client)["reference_column"]
    assert after["column"]["reference"]["symbol"] == "QQQ" and after["column"]["status"] == "done"
    assert after["pending"] is None


def test_switching_in_the_middle_cancels_the_previous_job(ctx, db):
    client = ctx["client"]
    _choose(client, "SPY")
    spy_id = ctx["queued"][0]
    _choose(client, "QQQ")  # trocar antes do job de SPY rodar
    ctx["run"](spy_id)
    ctx["queued"].remove(spy_id)
    assert (
        db.query(BenchmarkComparison).filter(BenchmarkComparison.column_id == spy_id).count() == 0
    )
    ctx["run"]()
    assert _board(client)["reference_column"]["column"]["reference"]["symbol"] == "QQQ"


def test_expired_candidate_sample_is_without_comparison(db, source, monkeypatch):
    user = make_user(db, "carol", admin=True)
    _seed(db, actor=str(user.id), expired={"RS-SOL1123D4"})
    queued = []
    monkeypatch.setattr(refcol, "enqueue", lambda cid: queued.append(cid))
    refcol.set_reference(db, SWEEP, str(user.id), "SPY", enqueue_fn=queued.append, source=source)
    Session = sessionmaker(bind=db.get_bind())
    refcol.run_column(
        queued[0], db_factory=Session, source=source, regenerate=lambda _d, r: TRADES.get(r.id, [])
    )
    state = refcol.state(db, SWEEP, str(user.id))
    assert state["column"]["counts"][bench.SAMPLE_EXPIRED] == 1
    rows = [{"result_id": "RS-SOL1123D4", "eligibility": "eligible"}]
    refcol.attach_to_leaderboard(db, SWEEP, str(user.id), rows)
    cell = rows[0]["reference_comparison"]
    assert cell["reason_code"] == bench.SAMPLE_EXPIRED
    assert cell["reason"] == "Amostra do candidato vencida em 30/09/2026."
    sols = [c for c in state["candidates"] if c["result_id"] == "RS-SOL1123D4"]
    assert sols[0]["sample"] == "expired"


def test_column_pins_samples_while_active_and_releases_after(ctx, db):
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    assert pinned(db, "snap0c1123mix") is True
    _choose(client, None)
    db.expire_all()
    assert pinned(db, "snap0c1123mix") is False


def test_other_sweep_and_other_user_start_without_column(ctx, db, source):
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    _seed(db, sweep_id="other-sweep", actor=ctx["actor"])
    assert refcol.state(db, "other-sweep", ctx["actor"])["column"] is None
    assert refcol.state(db, SWEEP, "someone-else")["column"] is None


def test_marked_reference_is_refused(ctx, db):
    from app.models_asset_catalog import AssetCatalog

    columns = {c.name for c in AssetCatalog.__table__.columns}
    row = {
        "asset_id": "binance:spot:WAVES/USDT",
        "symbol": "WAVES/USDT",
        "asset_class": "crypto",
        "state": "removed",
    }
    extra = {}
    if "name" in columns:
        extra["name"] = "Waves"
    if "state_reason" in columns:
        extra["state_reason"] = "Saiu da lista spot da Binance em 17 jun. 2024."
    for col in AssetCatalog.__table__.columns:
        if (
            not col.nullable
            and col.name not in row
            and col.name not in extra
            and col.default is None
            and col.server_default is None
            and not col.primary_key
        ):
            extra[col.name] = (
                "x" if str(col.type).startswith("VARCHAR") else datetime.now(timezone.utc)
            )
    db.add(AssetCatalog(**row, **extra))
    db.commit()
    res = _choose(ctx["client"], "WAVES/USDT")
    assert res.status_code == 422
    assert "não pode ser escolhido como referência" in res.json()["detail"]["message"]


def test_export_with_extra_columns_all_pages_in_ranking_order(ctx):
    client = ctx["client"]
    base = client.get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard/export",
        params={"metric": "delta_cagr_vs_bh"},
    )
    wb = openpyxl.load_workbook(io.BytesIO(base.content))
    header = [c.value for c in wb["Ranking"][1]]
    assert header[18:40] == list(xc.CONTEXT_FIELDS)
    assert not any(str(h).endswith("(extra)") for h in header)
    _choose(client, "SPY")
    ctx["run"]()
    page = _board(client, limit=2)
    res = client.get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard/export",
        params={"metric": "delta_cagr_vs_bh"},
    )
    assert res.status_code == 200
    assert res.headers["X-Export-Extra"] == "included"
    sync_e2e_fixture("descoberta.xlsx", res.content)
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    sheet = wb["Ranking"]
    header = [c.value for c in sheet[1]]
    assert header[40:] == list(xc.EXTRA_FIELDS)
    rows = [dict(zip(header, [c.value for c in r])) for r in sheet.iter_rows(min_row=2)]
    full = _board(client, limit=200)
    assert len(rows) == full["total"] > len(page["results"])
    assert [r["Candidato"].split(" · ")[0] for r in rows] == [r["symbol"] for r in full["results"]]
    by = {r["Candidato"].split(" · ")[0]: r for r in rows}
    assert by["BTC/USDT"]["Referência"] == "Próprio ativo (buy-and-hold, long-only)"
    assert by["BTC/USDT"]["Referência (extra)"] == "SPY · ETF · USD"
    assert by["BTC/USDT"]["Aviso de moeda (extra)"] == bench.CURRENCY_NOTE_TEXT
    assert by["NVDA"]["Tipo de retorno da referência"] == "Retorno de preço, sem dividendos"
    assert by["NVDA"]["Aviso de moeda (extra)"] in ("", None)
    pengu = by["PENGU/USDT"]
    assert (
        pengu["Situação da comparação"]
        == "Histórico insuficiente: Período comum de 14 dias: o mínimo é 30."
    )
    assert (
        pengu["Situação da comparação (extra)"]
        == "Histórico insuficiente: Período comum de 14 dias: o mínimo é 30."
    )
    for name in (
        "Retorno da estratégia no período comum",
        "CAGR da referência",
        "Diferença de CAGR (p.p.)",
    ):
        assert pengu[name] in (None, "")
        assert pengu[f"{name} (extra)"] in (None, "")
    legend = [r[0].value for r in wb["Legenda"].iter_rows(min_row=2)]
    assert "(extra)" in legend and "Retornos do ranking" in legend
    # 22 campos do próprio ativo: retorno derivado do CAGR gravado e da janela.
    btc = by["BTC/USDT"]
    assert btc["CAGR da estratégia no período comum"] == pytest.approx(0.82)
    assert btc["Retorno da estratégia no período comum"] == pytest.approx(
        (1.82) ** (2191 / 365.25) - 1
    )
    assert btc["Situação da comparação"] == "Válida"
    assert btc["Fora do período comum"] == "nada"
    # "Δ vs B&H" sai em p.p., como foi gravado, com o formato de "Diferença de CAGR (p.p.)".
    assert btc["Δ vs B&H"] == pytest.approx((0.82 - 0.7057) * 100)
    assert by["NVDA"]["Δ vs B&H"] == pytest.approx((0.65 - 0.848) * 100)
    delta_col = header.index("Δ vs B&H") + 1
    pp_col = header.index("Diferença de CAGR (p.p.)") + 1
    btc_line = 2 + [r["Candidato"].split(" · ")[0] for r in rows].index("BTC/USDT")
    assert sheet.cell(btc_line, delta_col).number_format == xc.PP_FORMAT
    assert sheet.cell(btc_line, pp_col).number_format == xc.PP_FORMAT


def test_promotion_inherits_the_recorded_cell(ctx, db):
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    inherited = refcol.inherited_for_promotion(db, SWEEP, ctx["actor"], "RS-BTC1123A1")
    assert inherited is not None and inherited["reference"]["symbol"] == "SPY"
    assert refcol.inherited_for_promotion(db, SWEEP, ctx["actor"], "RS-PENGU1123E5") is None


def test_leaderboard_metric_is_still_only_calmar_or_delta(ctx):
    res = ctx["client"].get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard", params={"metric": "reference_column"}
    )
    assert res.status_code == 400


def test_state_lists_candidate_windows_for_the_choice_warning(ctx):
    state = ctx["client"].get(f"/api/combos/discovery/sweeps/{SWEEP}/reference-column").json()
    windows = {c["symbol"]: c for c in state["candidates"]}
    assert (
        windows["PENGU/USDT"]["start"] == "2024-12-17"
        and windows["PENGU/USDT"]["end"] == "2024-12-31"
    )
    assert "AAPL" not in windows and state["not_computed"] == 1
    assert state["min_days"] == 30


def test_discovery_cell_dates_and_summary_never_carry_numbers_while_pending(ctx):
    client = ctx["client"]
    _choose(client, "SPY")
    # Antes do job: nenhuma célula com número (a coluna ainda não existe).
    board = _board(client)
    assert all(r.get("reference_comparison") is None for r in board["results"])
    assert board["reference_column"]["pending"]["reference"]["symbol"] == "SPY"


def test_sweep_still_running_refuses_a_choice(db, source):
    from app.models_discovery import DiscoverySweep

    user = make_user(db, "dan", admin=True)
    _seed(db, actor=str(user.id))
    sweep = db.query(DiscoverySweep).get(SWEEP)
    sweep.state = "running"
    db.commit()
    client = client_for(db, [discovery_routes.router], {"user": str(user.id), "admin": True})
    assert _choose(client, "SPY").status_code == 409


def test_calculation_failure_keeps_recorded_cells_and_retry_resumes(ctx, db, source, monkeypatch):
    client = ctx["client"]
    _choose(client, "SPY")
    # Lotes de 2: o primeiro lote fica gravado antes de o worker cair no segundo.
    monkeypatch.setattr(refcol, "BATCH_SIZE", 2)
    original_side = refcol._candidate_side
    calls = {"n": 0}

    def crashing_side(result):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("worker caiu")
        return original_side(result)

    monkeypatch.setattr(refcol, "_candidate_side", crashing_side)
    Session = sessionmaker(bind=db.get_bind())
    column_id = ctx["queued"].pop(0)
    first_run: list[str] = []

    def regenerate_first(_db, result):
        first_run.append(result.id)
        return TRADES.get(result.id, [])

    refcol.run_column(column_id, db_factory=Session, source=source, regenerate=regenerate_first)
    db.expire_all()
    column = db.query(DiscoveryReferenceColumn).get(column_id)
    assert column.status == "failed" and "worker caiu" in column.error
    recorded = {
        row.result_id: row.id
        for row in db.query(BenchmarkComparison).filter(BenchmarkComparison.column_id == column_id)
    }
    assert len(recorded) == 2 and set(recorded) == set(first_run)

    res = client.post(f"/api/combos/discovery/sweeps/{SWEEP}/reference-column/retry")
    assert res.status_code == 202
    assert res.json()["pending"]["status"] == "computing"
    monkeypatch.setattr(refcol, "_candidate_side", original_side)
    resumed: list[str] = []

    def regenerate_resume(_db, result):
        resumed.append(result.id)
        return TRADES.get(result.id, [])

    refcol.run_column(
        ctx["queued"].pop(0), db_factory=Session, source=source, regenerate=regenerate_resume
    )
    db.expire_all()
    column = db.query(DiscoveryReferenceColumn).get(column_id)
    assert column.status == "done" and column.error is None
    assert column.done == column.total == 5
    after = {
        row.result_id: row.id
        for row in db.query(BenchmarkComparison).filter(BenchmarkComparison.column_id == column_id)
    }
    # Retoma do que falta: as células gravadas ficam (mesmo id) e só as que faltavam são calculadas.
    assert len(after) == 5
    assert {rid: after[rid] for rid in recorded} == recorded
    assert sorted(resumed) == sorted(set(after) - set(recorded))
    board = _board(client)
    assert board["reference_column"]["column"]["counts"]["valid"] == 4


def test_candidate_without_recorded_cagr_closes_without_number_when_done(ctx, db):
    """CAGR acima do teto (``cagr=None`` sem ser ``insufficient_sample``): fora do cálculo."""

    from app.models_discovery import DiscoveryResult

    db.add(
        DiscoveryResult(
            id="RS-DOGE1123G7",
            sweep_id=SWEEP,
            combination_id=60,
            template_id="multi_ma_crossover",
            symbol="DOGE/USDT",
            timeframe="1d",
            direction="long",
            parameters={"direction": "long", "fast": 10},
            start_at=datetime(2019, 1, 1),
            end_at=datetime(2024, 12, 31),
            metrics={"oos_verdict": {"status": "GO"}},
            trades_count=40,
            calmar_ratio=1.5,
            cagr=None,
            benchmark_cagr=0.5,
            delta_cagr_vs_bh=None,
            coverage=1.0,
            strategy_identity_key="id-RS-DOGE1123G7",
            evidence_fingerprint="fp-RS-DOGE1123G7",
            eligibility="eligible",
            dedup_state="unique",
            asset_class="crypto",
            data_source="Binance Spot",
            session="24x7",
            adjustment="na",
            snapshot_id="snap0c1123mix",
            candle_version="v1",
        )
    )
    db.commit()
    assert refcol.cell_summary(None, eligibility="eligible") == {"state": "pending"}
    client = ctx["client"]
    _choose(client, "SPY")
    ctx["run"]()
    assert "RS-DOGE1123G7" not in ctx["regenerated"]
    board = _board(client, limit=200)
    col = board["reference_column"]["column"]
    assert col["status"] == "done" and col["total"] == 5 and col["counts"]["valid"] == 4
    cell = _cells(board)["RS-DOGE1123G7"]
    assert cell == {
        "state": "insufficient",
        "id": None,
        "reason_code": refcol.NO_CAGR,
        "reason": refcol.NO_CAGR_SHORT,
        "expired_at": None,
    }
    assert _cells(board)["RS-AAPL1123F6"] == {"state": "no_sample"}
    res = client.get(
        f"/api/combos/discovery/sweeps/{SWEEP}/leaderboard/export",
        params={"metric": "delta_cagr_vs_bh"},
    )
    sheet = openpyxl.load_workbook(io.BytesIO(res.content))["Ranking"]
    header = [c.value for c in sheet[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in sheet.iter_rows(min_row=2)]
    by = {r["Candidato"].split(" · ")[0]: r for r in rows}
    doge = by["DOGE/USDT"]
    assert doge["Referência (extra)"] == "SPY · ETF · USD"
    assert doge["Situação da comparação (extra)"] == refcol.NO_CAGR_TEXT
    assert doge["Diferença de CAGR (p.p.) (extra)"] in (None, "")
    assert by["AAPL"]["Referência (extra)"] == "SPY · ETF · USD"
