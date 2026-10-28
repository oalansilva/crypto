"""Card #1123: planilhas de Favoritos e de trades com o mesmo contexto (openpyxl)."""

from __future__ import annotations

import io

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import FavoriteStrategy
from app.routes import benchmark_routes, combo_routes, favorites as favorites_route
from app.services import benchmark_comparison as bench
from app.services import export_context as xc
from support_card_1123 import (
    sync_e2e_fixture,
    BTC_TRADES,
    FakeSource,
    btc_result,
    client_for,
    ensure_snapshot_table,
    make_user,
)

NR = "não registrado"


@pytest.fixture
def db(postgres_isolation, unit_database_url):
    engine = create_engine(unit_database_url)
    Base.metadata.create_all(bind=engine)
    Session = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    session = Session()
    ensure_snapshot_table(session)
    try:
        yield session
    finally:
        session.close()
        engine.dispose()


@pytest.fixture
def api(db, monkeypatch):
    class Disabled:
        enabled = False

    monkeypatch.setattr(favorites_route, "_FAVORITE_OHLCV_REPO", Disabled())
    source = FakeSource()
    bench.set_reference_source(source)
    admin = make_user(db, "alan", admin=True)
    state = {"user": str(admin.id), "admin": True}
    client = client_for(
        db, [benchmark_routes.router, favorites_route.router, combo_routes.router], state
    )
    yield client, admin, source
    bench.set_reference_source(None)


def _sheet(content: bytes, name: str):
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb[name]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    return wb, header, rows


def _favorite(
    db,
    user,
    name,
    symbol,
    *,
    metrics=None,
    comparison_id=None,
    start="2023-01-01",
    end="2025-10-31",
    snapshot_id=None,
):
    row = FavoriteStrategy(
        user_id=str(user.id),
        name=name,
        symbol=symbol,
        timeframe="1d",
        strategy_name="multi_ma_crossover",
        parameters={"direction": "long", "fast": 10},
        metrics=metrics or {"total_trades": 3, "total_return": 0.25, "win_rate": 0.5},
        start_date=start,
        end_date=end,
        data_source="ccxt" if "/" in symbol else "binance_equity",
        session="24x7" if "/" in symbol else "regular_ny",
        adjustment="none" if "/" in symbol else "split",
        benchmark_comparison_id=comparison_id,
        snapshot_id=snapshot_id,
    )
    db.add(row)
    db.commit()
    return row


def _record(db, user, strategy_symbol, reference, *, trades=BTC_TRADES, asset=None, window=None):
    strategy = bench.strategy_side(
        symbol=strategy_symbol,
        asset=asset,
        window=window or {"start": "2019-01-02", "end": "2024-12-31"},
        trades=trades,
    )
    return bench.create_comparison(
        db, strategy, reference, owner_kind="favorite", user_id=str(user.id), trigger="choice"
    )


def test_favorites_spreadsheet_has_todays_columns_context_and_old_favorite_states(api, db):
    client, admin, _ = api
    nvda_asset = {"symbol": "NVDA", "asset_class": "stock", "currency": "USD"}
    nvda_cmp = _record(db, admin, "NVDA", "SPY", asset=nvda_asset, trades=[])
    spy_cmp = _record(
        db,
        admin,
        "SPY",
        "BTC/USDT",
        asset={"symbol": "SPY", "asset_class": "etf", "currency": "USD"},
        trades=[],
        window={"start": "2015-01-02", "end": "2026-10-23"},
    )
    nvda = _favorite(
        db,
        admin,
        "NVDA Trend",
        "NVDA",
        comparison_id=nvda_cmp["id"],
        snapshot_id="snap-nvda-fav",
        metrics={
            "total_trades": 2,
            "analysis_asset": nvda_asset,
            "analysis_execution": {
                "precision_requested": "intrabar",
                "precision_effective": "candle",
                "window": {
                    "covered": {
                        "start": "2014-11-11T05:00:00Z",
                        "end": "2026-10-23T04:00:00Z",
                        "candles": 2994,
                    }
                },
            },
            "data_provenance": {
                "data_source": "Binance Ações",
                "session": "rth",
                "adjustment": "adjusted",
                "version_label": "3",
            },
        },
    )
    spy = _favorite(
        db, admin, "SPY Trend", "SPY", comparison_id=spy_cmp["id"], snapshot_id="snap-spy"
    )
    old = _favorite(db, admin, "BTC antigo", "BTC/USDT")
    res = client.get("/api/favorites/export", params={"ids": f"{old.id},{nvda.id},{spy.id}"})
    assert res.status_code == 200, res.text
    wb, header, rows = _sheet(res.content, "Favoritos")
    sync_e2e_fixture("favoritos.xlsx", res.content)
    assert header[:2] == ["Name", "Symbol"]
    assert header[24:] == list(xc.CONTEXT_FIELDS)
    assert [r["Name"] for r in rows] == ["BTC antigo", "NVDA Trend", "SPY Trend"]
    old_row, nvda_row, spy_row = rows
    assert old_row["Classe"] == "Cripto (deduzido)" and old_row["Moeda"] == "USDT (deduzido)"
    for field in (
        "Fonte",
        "Sessão",
        "Ajuste",
        "Versão da amostra",
        "Situação da comparação",
        "Versão da amostra da referência",
        "Comparação gravada em",
        "Período efetivo",
        "Modo de execução",
    ):
        assert old_row[field] == NR, field
    assert old_row["Referência"] == "Próprio ativo"
    for field in (
        "Retorno da estratégia no período comum",
        "CAGR da referência",
        "Diferença de CAGR (p.p.)",
    ):
        assert old_row[field] in (None, "")
    assert old_row["Período pedido"] == "01/01/2023 – 31/10/2025"
    assert nvda_row["Classe"] == "Ação" and nvda_row["Fonte"] == "Binance Ações"
    assert (
        nvda_row["Sessão"] == "Pregão regular de Nova York"
        and nvda_row["Ajuste"] == "Ajustado por split"
    )
    assert nvda_row["Referência"] == "SPY · ETF · USD"
    assert nvda_row["Tipo de retorno da referência"] == "Retorno de preço, sem dividendos"
    assert nvda_row["Aviso de moeda"] in (None, "")
    assert nvda_row["Situação da comparação"] == "Válida"
    assert nvda_row["Período efetivo"] == "11/11/2014 – 23/10/2026 · 2.994 candles"
    assert spy_row["Referência"] == "BTC/USDT · Cripto · USDT"
    assert spy_row["Tipo de retorno da referência"] == "Retorno de preço"
    assert spy_row["Aviso de moeda"] == bench.CURRENCY_NOTE_TEXT
    assert isinstance(spy_row["CAGR da referência"], float)
    legend = {r[0].value: r[1].value for r in wb["Legenda"].iter_rows(min_row=2)}
    assert "Tela × planilha" in legend and "Comparação refeita" in legend
    assert "deduzido" in legend and "não registrado" in legend
    assert "USDT tratado como USD" in legend and "Retorno de preço, sem dividendos" in legend


def test_favorites_export_is_admin_only(api, db):
    _, admin, _ = api
    common = client_for(db, [favorites_route.router], {"user": str(admin.id), "admin": False})
    assert common.get("/api/favorites/export").status_code == 403


def test_trades_spreadsheet_has_trades_context_and_legend(api, db):
    client, admin, _ = api
    created = client.post(
        "/api/benchmark-comparisons",
        json={"owner": "result", "reference_symbol": "SPY", "context": btc_result()},
    ).json()
    body = {
        "trades": BTC_TRADES,
        "symbol": "BTC/USDT",
        "template_name": "multi_ma_crossover",
        "timeframe": "1d",
        "benchmark_comparison_id": created["id"],
        "context": {**btc_result(), "is_optimization": True, "data_provenance": None},
    }
    res = client.post("/api/combos/export-trades", json=body)
    assert res.status_code == 200, res.text
    wb, header, rows = _sheet(res.content, "Trades")
    sync_e2e_fixture("trades.xlsx", res.content)
    assert wb.sheetnames == ["Trades", "Contexto", "Legenda"]
    assert header[-1] == "Moeda" and rows[0]["Moeda"] == "USDT"
    ctx = {r[0].value: (r[1].value, r[2].value) for r in wb["Contexto"].iter_rows(min_row=2)}
    assert list(ctx.keys()) == list(xc.CONTEXT_FIELDS)
    assert ctx["Referência"][0] == "SPY · ETF · USD"
    assert ctx["Período comum"][0] == "02/01/2019 – 31/12/2024 · 2.190 dias"
    assert ctx["Fora do período comum"][0] == "01/01/2019 (sem pregão em Nova York) · 0 trades"
    assert ctx["Aviso de moeda"][0] == bench.CURRENCY_NOTE_TEXT
    assert ctx["Modo de execução"][0] == "Otimização · intrabar 15m"
    assert ctx["Classe"] == ("Cripto", "registrado")
    assert isinstance(ctx["Diferença de CAGR (p.p.)"][0], float)
    assert wb["Contexto"]["B15"].number_format == xc.PCT_FORMAT  # CAGR da estratégia como número %


def test_trades_spreadsheet_with_insufficient_reference_has_empty_numbers(api, db):
    client, admin, source = api
    created = client.post(
        "/api/benchmark-comparisons",
        json={"owner": "result", "reference_symbol": "PENGU/USDT", "context": btc_result()},
    ).json()
    assert created["status"] == "insufficient"
    res = client.post(
        "/api/combos/export-trades",
        json={
            "trades": BTC_TRADES,
            "symbol": "BTC/USDT",
            "template_name": "t",
            "timeframe": "1d",
            "benchmark_comparison_id": created["id"],
            "context": btc_result(),
        },
    )
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ctx = {r[0].value: r[1].value for r in wb["Contexto"].iter_rows(min_row=2)}
    assert ctx["Situação da comparação"].startswith(
        "Histórico insuficiente: Período comum de 14 dias: o mínimo é 30."
    )
    for field in (
        "Retorno da estratégia no período comum",
        "CAGR da estratégia no período comum",
        "Retorno da referência",
        "CAGR da referência",
        "Diferença de CAGR (p.p.)",
    ):
        assert ctx[field] in (None, ""), field
    assert ctx["Aviso de moeda"] in (None, "")
    assert ctx["Tipo de retorno da referência"] == "Retorno de preço"


def test_trades_spreadsheet_from_a_favorite_reads_the_favorite_context(api, db):
    client, admin, _ = api
    cmp_ = _record(
        db,
        admin,
        "BTC/USDT",
        "SPY",
        asset={"symbol": "BTC/USDT", "asset_class": "crypto", "currency": "USDT"},
    )
    fav = _favorite(db, admin, "BTC antigo", "BTC/USDT", comparison_id=cmp_["id"])
    res = client.post(
        "/api/combos/export-trades",
        json={
            "trades": BTC_TRADES,
            "symbol": "BTC/USDT",
            "template_name": "t",
            "timeframe": "1d",
            "favorite_id": fav.id,
            "context": {"asset": {"asset_class": "stock"}},
        },
    )
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ctx = {r[0].value: (r[1].value, r[2].value) for r in wb["Contexto"].iter_rows(min_row=2)}
    assert ctx["Classe"] == ("Cripto (deduzido)", "deduzido")  # do favorito, não do corpo
    assert ctx["Fonte"] == (NR, NR)
    assert ctx["Referência"][0] == "SPY · ETF · USD"
    legend = [r[0].value for r in wb["Legenda"].iter_rows(min_row=2)]
    assert "Comparação refeita" in legend


def test_the_three_workbooks_share_the_same_context_headers():
    assert len(xc.CONTEXT_FIELDS) == 22
    assert xc.CONTEXT_FIELDS[0] == "Classe" and xc.CONTEXT_FIELDS[-1] == "Comparação gravada em"
    assert len(xc.EXTRA_FIELDS) == 13
    assert xc.EXTRA_FIELDS[0] == "Referência (extra)"
    assert xc.EXTRA_FIELDS[-1] == "Comparação gravada em (extra)"


def test_cells_never_turn_insufficient_into_zero():
    cells = xc.comparison_cells(
        {
            "status": "insufficient",
            "reference": {"symbol": "CRCL", "asset_class": "stock"},
            "reason": {"code": "no_overlap", "text": "x", "short": "y"},
        }
    )
    numeric = cells[4:9]
    assert all(c.rendered() in ("", None) for c in numeric)
    assert cells[9].rendered() == "Histórico insuficiente: x"
